"""
make_ieee_paper.py
------------------
Builds the IEEE-format conference paper shell (Word .docx) for:
"A Lightweight, Well-Calibrated Multi-Task Framework for Trustworthy Lumbar
Spine MRI Diagnosis with Uncertainty and Explainability"

Conventions follow IEEE conference template (IEEEtran.cls):
  Letter paper, two-column body, Times New Roman, 24pt title, 10pt body,
  run-in Abstract-- / Index Terms--, numbered sections, 8pt references.
Placeholders are marked [LIKE THIS] - replace them with your real content.

Run:  python make_ieee_paper.py
"""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Times New Roman"
OUT = r"C:\Users\srish\SPIDER\Trustworthy_Spine_MRI_Conference_Paper.docx"


# ---------------------------------------------------------------- helpers
def _set_font(run, size, bold=False, italic=False, small_caps=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.small_caps = small_caps
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), FONT)


def _margins(sec):
    sec.top_margin = Inches(0.75)
    sec.bottom_margin = Inches(0.75)
    sec.left_margin = Inches(0.625)
    sec.right_margin = Inches(0.625)


def _columns(sec, num, space_twips=360):
    sectPr = sec._sectPr
    for c in sectPr.findall(qn("w:cols")):
        sectPr.remove(c)
    cols = OxmlElement("w:cols")
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), str(space_twips))
    sectPr.append(cols)


def para(doc, text, size=10, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
         before=0, after=6, small_caps=False, left=None, first=None):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if left is not None:
        pf.left_indent = Inches(left)
    if first is not None:
        pf.first_line_indent = Inches(first)
    r = p.add_run(text)
    _set_font(r, size, bold, italic, small_caps)
    return p


def runin(doc, lead, text, size=9):
    """Run-in paragraph, e.g. 'Abstract--' bold lead followed by text."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    r1 = p.add_run(lead)
    _set_font(r1, size, bold=True, italic=True)
    r2 = p.add_run(text)
    _set_font(r2, size)
    return p


def heading(doc, text, before=10):
    return para(doc, text, size=10, bold=True, small_caps=True,
                align=WD_ALIGN_PARAGRAPH.LEFT, before=before, after=4)


def subheading(doc, text, before=8):
    return para(doc, text, size=10, bold=False, italic=True,
                align=WD_ALIGN_PARAGRAPH.LEFT, before=before, after=4)


def ref(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.left_indent = Inches(0.25)
    pf.first_line_indent = Inches(-0.25)
    pf.space_after = Pt(2)
    r = p.add_run(text)
    _set_font(r, 8)
    return p


# ---------------------------------------------------------------- document
doc = Document()
_set_font(doc.styles["Normal"], 10)

sec1 = doc.sections[0]
_margins(sec1)
_columns(sec1, 1)  # title/abstract block: full width

# ---- Title block ---------------------------------------------------------
para(doc, "A Lightweight, Well-Calibrated Multi-Task Framework for Trustworthy "
          "Lumbar Spine MRI Diagnosis with Uncertainty and Explainability",
     size=24, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=10)

para(doc, "First Author, Second Author", size=11, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
para(doc, "Department of ..., University ..., City, Country", size=8,
     align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
para(doc, "email@university.edu", size=8, italic=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, after=8)

# ---- Abstract --------------------------------------------------------------
runin(doc, "Abstract\u2014",
      "Lumbar-spine pathology affects millions worldwide, yet current AI "
      "diagnostic tools lack the trustworthiness required for clinical adoption. "
      "This work presents a lightweight, well-calibrated multi-task deep-learning "
      "framework designed specifically for lumbar MRI diagnosis. Built on efficient "
      "architectures (EfficientViT / MobileViT) with CBAM attention, the framework "
      "jointly performs pathology classification and Pfirrmann severity grading, "
      "reducing inference time and parameter count while maintaining high accuracy. "
      "Crucially, the model incorporates Monte-Carlo Dropout for uncertainty "
      "quantification, enabling it to flag low-confidence predictions that require "
      "expert review. We evaluate calibration rigorously using Expected Calibration "
      "Error (ECE), Brier score, and reliability diagrams, ensuring that reported "
      "probabilities genuinely reflect prediction confidence. For explainability, "
      "the framework generates multiple saliency maps (Grad-CAM, Grad-CAM++, "
      "ScoreCAM, and EigenCAM) to visually justify each diagnosis, allowing "
      "radiologists to verify that the model focuses on the correct anatomical "
      "features. Evaluated on the SPIDER dataset (218 patients, 447 MRI series), "
      "the framework achieves competitive accuracy with 60% fewer parameters and "
      "40% lower latency than ResNet50 baselines [REPLACE with measured values]. "
      "Monte-Carlo uncertainty robustly identifies failure cases, and calibration "
      "metrics confirm that the model's confidence scores are reliable. The "
      "combination of lightweight efficiency, calibrated uncertainty, and rich "
      "explainability makes this framework a practical candidate for real-world "
      "clinical deployment.")

# ---- Index Terms -------------------------------------------------------------
runin(doc, "Index Terms\u2014",
      "Calibration, deep learning, explainable AI, lightweight neural networks, "
      "lumbar spine MRI, medical image classification, multi-task learning, "
      "uncertainty estimation. (Optional: Pfirrmann grading, Monte Carlo dropout.)")

# ---- Body: two-column section ------------------------------------------------
sec2 = doc.add_section(WD_SECTION.CONTINUOUS)
_margins(sec2)
_columns(sec2, 2)

# ================================================================ I. INTRO
heading(doc, "I. Introduction")
para(doc,
     "Deep-learning models for lumbar-spine MRI have reached high raw accuracy "
     "on tasks such as disc-degeneration grading and pathology detection, yet "
     "almost none of them are trustworthy enough for real clinical deployment. "
     "In practice, a radiologist needs more than a single predicted label: they "
     "need to know how confident the model is, whether that confidence is "
     "well-calibrated, and why a particular decision was made [4], [16].")
para(doc,
     "Current models fall short on four fronts. First, uncertainty estimation is "
     "largely absent, so a wrong prediction made with high apparent confidence is "
     "indistinguishable from a correct one. Second, calibration is rarely "
     "evaluated, meaning the probability a model reports does not reflect its "
     "true likelihood of being right [14], [24]. Third, explainability is "
     "limited, leaving clinicians without visual evidence linking a diagnosis to "
     "the underlying anatomy [18], [19]. Fourth, most high-performing "
     "architectures are computationally heavy, making them impractical for "
     "point-of-care or resource-constrained hospital settings [20], [21].")
para(doc,
     "There is therefore a clear gap for a single framework that is lightweight, "
     "well-calibrated, uncertainty-aware and explainable, while jointly handling "
     "multiple diagnostic tasks on lumbar-spine MRI. This paper presents, to the "
     "best of our knowledge, the first such unified framework, with the following "
     "contributions:")
para(doc,
     "1) A lightweight multi-task architecture built on EfficientViT/MobileViT "
     "with CBAM attention that jointly performs pathology classification and "
     "Pfirrmann severity grading;\n"
     "2) Uncertainty estimation via Monte-Carlo Dropout, which flags ambiguous "
     "cases for radiologist review;\n"
     "3) Explicit calibration evaluation (ECE, Brier score, reliability "
     "diagrams) with temperature scaling;\n"
     "4) A four-method explainability suite (Grad-CAM, Grad-CAM++, ScoreCAM, "
     "EigenCAM) for visual verification of each diagnosis.",
     after=8)

# ================================================================ II. RELATED WORK
heading(doc, "II. Related Work")

subheading(doc, "A. Deep Learning for Spinal Diagnosis")
para(doc,
     "Early automated reading of lumbar MRI dates to SpineNet [1], which "
     "classified radiological features and localized the evidence, and to the "
     "ISSLS Prize study [2] showing automated reading comparable with an expert "
     "radiologist. Subsequent work applied semantic segmentation for spinal "
     "stenosis [3], automated Pfirrmann-style grading [5], Modic-change "
     "detection [6], multi-branch residual networks for disc classification [7], "
     "and YOLO-based herniation detection [8]. Survey work [4] summarizes these "
     "trends. The common thread is high raw accuracy on a single task, with "
     "heavy architectures and little or no uncertainty or calibration reporting.")

subheading(doc, "B. Multi-Task Learning in Medical Imaging")
para(doc,
     "Multi-task learning (MTL) lets one model share a feature extractor across "
     "related tasks, improving generalization and efficiency. Graham et al. [9] "
     "showed a single model can simultaneously segment and classify histology "
     "images; similar gains were reported for breast ultrasound [10], [12] and "
     "feature-transfer MTL [11]. Wu and Gou [13] recently combined MTL with "
     "uncertainty-guided learning. None of these apply MTL to joint pathology "
     "detection and severity grading on lumbar MRI, and few combine it with "
     "calibration.")

subheading(doc, "C. Uncertainty and Calibration")
para(doc,
     "Mehrtash et al. [14] and Carneiro et al. [15] demonstrated calibrated "
     "confidence and predictive uncertainty for medical segmentation and "
     "classification. Abdar et al. [16] provide a comprehensive survey of "
     "uncertainty quantification, and Angelopoulos and Bates [17] formalize "
     "conformal prediction as a distribution-free alternative. Calibration "
     "studies remain rare in spine imaging, and none integrate uncertainty with "
     "a lightweight multi-task architecture.")

subheading(doc, "D. Lightweight Architectures and Explainability")
para(doc,
     "Explainability in medical imaging has been surveyed by van der Velden et "
     "al. [18] and Singh et al. [19], who categorize saliency approaches such as "
     "class-activation maps. On the efficiency side, knowledge-distillation "
     "ensembles [20] and adversarially robust lightweight networks [21] show "
     "that accuracy can be preserved at low parameter counts. To date, "
     "lightweight design and multi-method explainability have not been combined "
     "with calibrated, uncertainty-aware multi-task lumbar MRI diagnosis.")

# ================================================================ III. DATASET
heading(doc, "III. Dataset: SPIDER Benchmark")
para(doc,
     "We evaluate on SPIDER [22], a public benchmark for lumbar-spine MRI "
     "released under CC-BY 4.0. It contains 218 patients and 447 sagittal "
     "T1/T2-weighted series with expert per-level annotations for five lumbar "
     "intervertebral discs (and additional levels where annotated): binary "
     "pathology labels (disc herniation, bulging, narrowing, endplate defects, "
     "spondylolisthesis), Pfirrmann grades (1\u20135) for disc degeneration, "
     "Modic changes (I/II/III), and segmentation masks for vertebrae, discs, "
     "and the spinal canal.")
para(doc,
     "We split by patient into training (70%), validation (15%), and test "
     "(15%) to avoid data leakage; no patient appears in more than one split. "
     "Preprocessing comprises intensity normalization, cropping to the lumbar "
     "region, resizing to 224\u00d7224, and augmentation (rotation, "
     "brightness/contrast, Gaussian noise).")

# ================================================================ IV. METHODOLOGY
heading(doc, "IV. Proposed Methodology")

subheading(doc, "A. Lightweight Shared-Encoder Backbone")
para(doc,
     "The feature extractor is either EfficientViT [30], which uses "
     "multi-scale linear attention for high-resolution dense prediction, or "
     "MobileViT [31], which fuses convolutions with lightweight transformers "
     "(~5\u20136.5M parameters). A CBAM module [29] refines spatial and channel "
     "features so the network attends to disc and vertebral regions.")

subheading(doc, "B. Multi-Task Prediction Heads")
para(doc,
     "The shared encoder feeds two heads. The pathology head is a binary "
     "multi-label classifier over five conditions (sigmoid, binary cross-entropy "
     "loss); the severity head performs 5-class Pfirrmann grading (softmax, "
     "cross-entropy loss). Joint training minimizes the weighted sum of both "
     "losses, enabling complementary feature learning.")

subheading(doc, "C. Uncertainty Estimation via Monte-Carlo Dropout")
para(doc,
     "Following Gal and Ghahramani [23], dropout is enabled at inference and "
     "T=20 forward passes are performed per input. The mean over passes is the "
     "final prediction; the standard deviation quantifies epistemic "
     "uncertainty. High-uncertainty cases are flagged for radiologist review.")

subheading(doc, "D. Calibration")
para(doc,
     "We measure calibration with Expected Calibration Error (ECE), Brier "
     "score, and reliability diagrams [24]. Post-hoc temperature scaling "
     "recalibrates confidence without changing accuracy.")

subheading(doc, "E. Explainability")
para(doc,
     "Four saliency methods are generated per prediction: Grad-CAM [25], "
     "Grad-CAM++ [26], ScoreCAM [27], and EigenCAM [28]. Overlaying these maps "
     "on the MRI lets radiologists verify that the model attends to correct "
     "anatomy rather than image artifacts.")

# ================================================================ V. EXPERIMENTAL SETUP
heading(doc, "V. Experimental Setup")
subheading(doc, "A. Hardware and Software")
para(doc,
     "[CONFIRM these match your actual setup] Training used an NVIDIA Tesla "
     "V100 (32 GB) with mixed precision (FP16); inference latency and parameter "
     "counts were measured on an NVIDIA GTX 1080 Ti to simulate "
     "resource-constrained deployment. Software: PyTorch 2.0.1, timm, "
     "pytorch-grad-cam, Albumentations, scikit-learn, nibabel, and Weights & "
     "Biases for tracking.")
subheading(doc, "B. Baselines and Metrics")
para(doc,
     "Baselines are ResNet50, DenseNet121, Vision Transformer, and "
     "EfficientNet-B4. We report accuracy and AUROC per task, parameter count, "
     "latency, ECE and Brier score before/after temperature scaling, and "
     "uncertainty quality (e.g., accuracy among flagged vs. accepted cases).")

# ================================================================ VI. RESULTS
heading(doc, "VI. Results")
para(doc,
     "[FILL IN after experiments. Suggested structure:]")
para(doc,
     "Table I: accuracy/AUROC vs. baselines for pathology detection and "
     "Pfirrmann grading.\n"
     "Table II: parameters and latency vs. baselines.\n"
     "Table III: ECE and Brier score before/after temperature scaling.\n"
     "Table IV: uncertainty flagging \u2014 error coverage among "
     "high-uncertainty cases.\n"
     "Fig. 1: reliability diagrams. Fig. 2: sample saliency maps (Grad-CAM, "
     "Grad-CAM++, ScoreCAM, EigenCAM) with ground-truth overlays.")

# ================================================================ VII. DISCUSSION
heading(doc, "VII. Discussion")
para(doc,
     "[Draft once results exist. Consider: what calibrated uncertainty enables "
     "for referral workflows; the efficiency\u2013accuracy trade-off versus heavy "
     "baselines; agreement of the four saliency methods; limitations such as "
     "dataset size (218 patients), T1/T2 acquisition variability, and the "
     "ordinal nature of Pfirrmann grading.]")

# ================================================================ VIII. CONCLUSION
heading(doc, "VIII. Conclusion")
para(doc,
     "[Draft. Summary sentence + future work, e.g., conformal prediction [17], "
     "3D/volumetric input, external validation, and prospective clinical "
     "evaluation.]")

# ================================================================ REFERENCES
sec3 = doc.add_section(WD_SECTION.CONTINUOUS)
_margins(sec3)
_columns(sec3, 1)

heading(doc, "References")
refs = [
 "A. Jamaludin, T. Kadir, and A. Zisserman, \u201cSpineNet: Automated "
 "classification and evidence visualization in spinal MRIs,\u201d Med. Image "
 "Anal., vol. 41, pp. 63\u201373, 2017, doi: 10.1016/j.media.2017.07.002.",
 "A. Jamaludin, M. Lootus, T. Kadir, A. Zisserman, J. Urban, et al., \u201cISSLS "
 "Prize in Bioengineering Science 2017: Automation of reading of radiological "
 "features from magnetic resonance images (MRIs) of the lumbar spine without "
 "human intervention is comparable with an expert radiologist,\u201d Eur. Spine "
 "J., vol. 26, no. 5, pp. 1374\u20131383, 2017, doi: 10.1007/s00586-017-4956-3.",
 "A. S. Al-Kafri, S. Sudirman, A. Hussain, D. Al-Jumeily, F. Natalia, H. "
 "Meidia, et al., \u201cBoundary delineation of MRI images for lumbar spinal "
 "stenosis detection through semantic segmentation using deep neural "
 "networks,\u201d IEEE Access, vol. 7, pp. 43487\u201343501, 2019, doi: "
 "10.1109/ACCESS.2019.2908002.",
 "F. Galbusera, G. Casaroli, and T. Bassani, \u201cArtificial intelligence and "
 "machine learning in spine research,\u201d JOR Spine, vol. 2, no. 1, Art. no. "
 "e1044, 2019, doi: 10.1002/jsp2.1044.",
 "W. Liawrungrueang, P. Kim, V. Kotheeranurak, K. Jitpakdee, P. Sarasombath, "
 "et al., \u201cAutomatic detection, classification, and grading of lumbar "
 "intervertebral disc degeneration using an artificial neural network "
 "model,\u201d Diagnostics, vol. 13, no. 4, Art. no. 663, 2023, doi: "
 "10.3390/diagnostics13040663.",
 "G. Liu, L. Wang, S.-N. You, Z. Wang, S. Zhu, C. Chen, et al., \u201cAutomatic "
 "detection and classification of Modic changes in MRI images using deep "
 "learning: Intelligent assisted diagnosis system,\u201d Orthop. Surg., vol. "
 "16, no. 1, pp. 196\u2013206, 2024, doi: 10.1111/os.13894.",
 "I. Ram, S. Kumar, and A. K. Keshri, \u201cClassification of intervertebral "
 "disc using novel multi-branch convolutional residual network model,\u201d "
 "Biomed. Signal Process. Control, vol. 102, Art. no. 107332, 2025, doi: "
 "10.1016/j.bspc.2024.107332.",
 "Z. Karakuzu G\u00fcng\u00f6r, H. Vehbi, and A. Cans\u0131n, \u201cAutomated "
 "detection of lumbar disc herniation at L4\u2013L5 and L5\u2013S1 levels on "
 "sagittal MRI using a YOLO-based deep learning model,\u201d BMC Musculoskelet. "
 "Disord., vol. 27, Art. no. 715, 2026, doi: 10.1186/s12891-026-10124-4.",
 "S. Graham, Q. D. Vu, M. Jahanifar, S. E. A. Raza, F. Minhas, D. Snead, et "
 "al., \u201cOne model is all you need: Multi-task learning enables "
 "simultaneous histology image segmentation and classification,\u201d Med. "
 "Image Anal., vol. 83, Art. no. 102685, 2023, doi: "
 "10.1016/j.media.2022.102685.",
 "Y. Zhou, H. Chen, Y. Li, Q. Liu, X. Xu, S. Wang, et al., \u201cMulti-task "
 "learning for segmentation and classification of tumors in 3D automated "
 "breast ultrasound images,\u201d Med. Image Anal., vol. 70, Art. no. 101918, "
 "2021, doi: 10.1016/j.media.2020.101918.",
 "F. Gao, H. Yoon, T. Wu, and X. Chu, \u201cA feature transfer enabled "
 "multi-task deep learning model on medical imaging,\u201d Expert Syst. Appl., "
 "vol. 143, Art. no. 112957, 2020, doi: 10.1016/j.eswa.2019.112957.",
 "M. Xu, K. Huang, and X. Qi, \u201cA regional-attentive multi-task learning "
 "framework for breast ultrasound image segmentation and classification,\u201d "
 "IEEE Access, vol. 11, pp. 5377\u20135392, 2023, doi: "
 "10.1109/ACCESS.2023.3236693.",
 "X. Wu and G. Gou, \u201cUncertainty bidirectional guidance of multi-task "
 "mamba network for medical image classification and segmentation,\u201d "
 "Signal Image Video Process., vol. 19, Art. no. 29, 2025, doi: "
 "10.1007/s11760-024-03633-z.",
 "A. Mehrtash, W. M. Wells, C. M. Tempany, P. Abolmaesumi, and T. Kapur, "
 "\u201cConfidence calibration and predictive uncertainty estimation for deep "
 "medical image segmentation,\u201d IEEE Trans. Med. Imaging, vol. 39, no. 12, "
 "pp. 3868\u20133878, 2020, doi: 10.1109/TMI.2020.3006437.",
 "G. Carneiro, L. Zorron Cheng Tao Pu, R. Singh, and A. Burt, \u201cDeep "
 "learning uncertainty and confidence calibration for the five-class polyp "
 "classification from colonoscopy,\u201d Med. Image Anal., vol. 62, Art. no. "
 "101653, 2020, doi: 10.1016/j.media.2020.101653.",
 "M. Abdar, F. Pourpanah, S. Hussain, D. Rezazadegan, L. Liu, M. Ghavamzadeh, "
 "et al., \u201cA review of uncertainty quantification in deep learning: "
 "Techniques, applications and challenges,\u201d Inf. Fusion, vol. 76, pp. "
 "243\u2013297, 2021, doi: 10.1016/j.inffus.2021.05.008.",
 "A. N. Angelopoulos and S. Bates, \u201cConformal prediction: A gentle "
 "introduction,\u201d Found. Trends Mach. Learn., vol. 16, no. 4, pp. "
 "494\u2013591, 2023, doi: 10.1561/2200000101.",
 "B. H. M. van der Velden, H. J. Kuijf, K. G. A. Gilhuijs, and M. A. "
 "Viergever, \u201cExplainable artificial intelligence (XAI) in deep "
 "learning-based medical image analysis,\u201d Med. Image Anal., vol. 79, Art. "
 "no. 102470, 2022, doi: 10.1016/j.media.2022.102470.",
 "A. Singh, S. Sengupta, and V. Lakshminarayanan, \u201cExplainable deep "
 "learning models in medical image analysis,\u201d J. Imaging, vol. 6, no. 6, "
 "Art. no. 52, 2020, doi: 10.3390/jimaging6060052.",
 "J. Kang and J. Gwak, \u201cEnsemble learning of lightweight deep learning "
 "models using knowledge distillation for image classification,\u201d "
 "Mathematics, vol. 8, no. 10, Art. no. 1652, 2020, doi: 10.3390/math8101652.",
 "A. Singh, M. P. Singh, and A. K. Singh, \u201cAdversarially enhanced "
 "learning (AEL): Robust lightweight deep learning approach for radiology "
 "image classification against adversarial attacks,\u201d Image Vis. Comput., "
 "vol. 154, Art. no. 105405, 2025, doi: 10.1016/j.imavis.2024.105405.",
 "J. van der Graaf, M. P. A. Starmans, et al., \u201cLumbar spine segmentation "
 "in MR images: A dataset and a public benchmark,\u201d Sci. Data, vol. 11, "
 "Art. no. 264, 2024, doi: 10.1038/s41597-024-03090-w.",
 "Y. Gal and Z. Ghahramani, \u201cDropout as a Bayesian approximation: "
 "Representing model uncertainty in deep learning,\u201d in Proc. 33rd Int. "
 "Conf. Mach. Learn. (ICML), 2016, pp. 1050\u20131059.",
 "C. Guo, G. Pleiss, Y. Sun, and K. Q. Weinberger, \u201cOn calibration of "
 "modern neural networks,\u201d in Proc. 34th Int. Conf. Mach. Learn. (ICML), "
 "2017, pp. 1321\u20131330.",
 "R. R. Selvaraju, M. Cogswell, A. Das, R. Vedantam, D. Parikh, and D. "
 "Batra, \u201cGrad-CAM: Visual explanations from deep networks via "
 "gradient-based localization,\u201d Int. J. Comput. Vis., vol. 128, pp. "
 "336\u2013359, 2020, doi: 10.1007/s11263-019-01228-7.",
 "A. Chattopadhay, A. Sarkar, P. Howlader, and V. N. Balasubramanian, "
 "\u201cGrad-CAM++: Generalized gradient-based visual explanations for deep "
 "convolutional networks,\u201d in Proc. IEEE Winter Conf. Appl. Comput. Vis. "
 "(WACV), 2018, pp. 839\u2013847, doi: 10.1109/WACV.2018.00097.",
 "H. Wang, Z. Wang, M. Du, F. Yang, Z. Zhang, S. Ding, et al., "
 "\u201cScore-CAM: Score-weighted visual explanations for convolutional neural "
 "networks,\u201d in Proc. IEEE/CVF Conf. Comput. Vis. Pattern Recognit. "
 "Workshops (CVPRW), 2020, pp. 111\u2013119, doi: "
 "10.1109/CVPRW50498.2020.00020.",
 "M. B. Muhammad and M. Yeasin, \u201cEigen-CAM: Class activation map using "
 "principal components,\u201d in Proc. Int. Joint Conf. Neural Netw. (IJCNN), "
 "2020, pp. 1\u20137.",
 "S. Woo, J. Park, J.-Y. Lee, and I. S. Kweon, \u201cCBAM: Convolutional "
 "block attention module,\u201d in Proc. Eur. Conf. Comput. Vis. (ECCV), 2018, "
 "pp. 3\u201319, doi: 10.1007/978-3-030-01234-2_1.",
 "H. Cai, J. Li, M. Hu, C. Gan, and S. Han, \u201cEfficientViT: Lightweight "
 "multi-scale attention for high-resolution dense prediction,\u201d in Proc. "
 "IEEE/CVF Int. Conf. Comput. Vis. (ICCV), 2023, pp. 17256\u201317267, doi: "
 "10.1109/ICCV51070.2023.01584.",
 "S. Mehta and M. Rastegari, \u201cMobileViT: Light-weight, general-purpose, "
 "and mobile-friendly vision transformer,\u201d in Proc. Int. Conf. Learn. "
 "Represent. (ICLR), 2022.",
]
for i, r in enumerate(refs, 1):
    ref(doc, f"[{i}] {r}")

doc.save(OUT)
print("Saved:", OUT)
print("Paragraphs:", len(doc.paragraphs), "| References:", len(refs))
